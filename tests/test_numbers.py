"""Dates, amounts and times: conversion, comparison, and choosing between
replies. Audit 2026-09-10: H-5, H-6, H-7, M-1, M-2, M-3, L-2, L-7, L-8, L-9,
L-10."""
import support  # noqa: F401  -- before trlib: sets the environment it reads

import unittest

import trlib
import trref

NBSP, WJ = "\u00a0", "\u2060"


class Localize(unittest.TestCase):
    """Segments that are only a value, converted without a model."""

    def test_swiss_identifiers_stay_verbatim(self):
        for v in ("12345", "80331", "123456789", "2024", "1250"):
            self.assertEqual(trlib.localize(v, "en", "de-CH"), v)

    def test_swiss_grouped_numbers_and_money(self):
        self.assertEqual(trlib.localize("12,450", "en", "de-CH"), f"12{NBSP}450")
        self.assertEqual(trlib.localize("CHF 12,450.50", "en", "de-CH"),
                         f"CHF 12{NBSP}450.50")
        self.assertEqual(trlib.localize("20.00 CHF", "en", "de-CH"), f"Fr.{NBSP}20.{WJ}–")

    def test_a_number_that_may_be_a_reference_stays(self):
        for tgt in ("de", "sl", "de-CH"):
            for v in ("5.10", "3.2", "14.30"):
                self.assertEqual(trlib.localize(v, "en", tgt), v, (v, tgt))

    def test_unambiguous_values_still_convert(self):
        self.assertEqual(trlib.localize("5.10 EUR", "en", "de"), "5,10 EUR")
        self.assertEqual(trlib.localize("0.75", "en", "de"), "0,75")
        self.assertEqual(trlib.localize("1,234.56", "en", "sl"), "1.234,56")
        self.assertEqual(trlib.localize("12.450,00 EUR", "sl", "en"), "12,450.00 EUR")
        self.assertEqual(trlib.localize("5. 3. 2024", "sl", "en"), "March 5, 2024")

    def test_a_value_keeps_what_stands_around_it(self):
        for v, s, t, want in [
                ("$1,000", "en", "de", "$1.000"),
                ("-1,250.00", "en", "de", "-1.250,00"),
                ("(1,250.00)", "en", "de", "(1.250,00)"),
                ("12.5%", "en", "de", "12,5%"),
                ("€12,450.00", "en", "de", "€12.450,00"),
                ("€ 12.450,00", "sl", "en", "€ 12,450.00"),
                ("-12.450,00", "sl", "en", "-12,450.00"),
                ("(CHF 1,250.00)", "en", "de-CH", f"(Fr.{NBSP}1250.{WJ}–)"),
                ("-12,450", "en", "de-CH", f"-12{NBSP}450"),
                ("12.5%", "en", "de-CH", "12,5%")]:
            self.assertEqual(trlib.localize(v, s, t), want, (v, s, t))

    def test_brackets_alone_do_not_make_a_quantity(self):
        self.assertEqual(trlib.localize("(3.2)", "en", "de"), "(3.2)")

    def test_ranges(self):
        self.assertEqual(trlib.localize("12:00–13:00", "sl", "en"),
                         "12:00 p.m.–1:00 p.m.")
        self.assertEqual(trlib.localize("12:00–13:00", "en", "de-CH"), "12.00–13.00")
        self.assertEqual(trlib.localize("2:30 p.m.–4:00 p.m.", "en", "de"), "14:30–16:00")
        self.assertEqual(trlib.localize("1,000-2,000", "en", "de"), "1.000-2.000")
        for v in ("2024-03-05", "5.10–5.12"):
            self.assertEqual(trlib.localize(v, "en", "de"), v)

    def test_slovene_month_names(self):
        for v in ("5. marec 2024", "5. marca 2024"):
            self.assertEqual(trlib.localize(v, "sl", "en"), "March 5, 2024")

    def test_every_currency_marks_an_amount(self):
        for v in ("12,450.00 GBP", "Fr. 20.00", "20.00 Fr.", "CHF 1250.50", "1.000 SIT"):
            self.assertFalse(trlib.is_translatable(v), v)


class FixNumericFormat(unittest.TestCase):
    """Numbers the model left in the source's format, inside sentences."""

    def fix(self, src, tgt, s, t):
        return trlib.fix_numeric_format(src, tgt, s, t)

    def test_references_and_times_are_not_decimals(self):
        for t in ("de", "sl", "de-CH"):
            for s in ("See Article 3.2 of the lease.", "See Section 5.10.",
                      "Articles 3.2 and 3.3 apply.", "The hearing is at 14.30.",
                      "It starts at 9.30 a.m.", "Under § 5.1 the fee is due."):
                self.assertEqual(self.fix(s, s, "en", t), s, (s, t))

    def test_a_swiss_time_in_the_draft_is_left(self):
        self.assertEqual(self.fix("The meeting starts at 10.30 am.",
                                  "Die Sitzung beginnt um 10.30 Uhr.", "en", "de-CH"),
                         "Die Sitzung beginnt um 10.30 Uhr.")
        self.assertEqual(self.fix("Arrive by 10.30.", "Kommen Sie bis 10.30 Uhr.",
                                  "en", "de-CH"), "Kommen Sie bis 10.30 Uhr.")

    def test_a_reference_word_in_the_draft_protects(self):
        self.assertEqual(self.fix("See 5.10 above.", "Siehe Abschnitt 5.10 oben.",
                                  "en", "de"), "Siehe Abschnitt 5.10 oben.")

    def test_amounts_still_convert(self):
        cases = [
            ("The fee is 12.50 EUR.", "Die Gebühr beträgt 12.50 EUR.", "en", "de",
             "Die Gebühr beträgt 12,50 EUR."),
            ("The amount is 12,450.00.", "Der Betrag ist 12,450.00.", "en", "de",
             "Der Betrag ist 12.450,00."),
            ("Znesek je 12.450,00.", "The amount is 12.450,00.", "sl", "en",
             "The amount is 12,450.00."),
            ("It was sold at 12.50 EUR.", "Es wurde zu 12.50 EUR verkauft.", "en", "de",
             "Es wurde zu 12,50 EUR verkauft."),
            ("The rate is 3.25 per cent.", "Der Satz ist 3.25 Prozent.", "en", "sl",
             "Der Satz ist 3,25 Prozent."),
        ]
        for src, tgt, s, t, want in cases:
            self.assertEqual(self.fix(src, tgt, s, t), want)

    def test_money_and_a_bare_number_in_one_swiss_sentence(self):
        self.assertEqual(self.fix("Interest of 2.50 on CHF 2.50",
                                  "Zins von 2.50 auf CHF 2.50", "en", "de-CH"),
                         "Zins von 2,50 auf CHF 2.50")

    def test_a_number_inside_a_longer_one_is_left(self):
        self.assertEqual(self.fix("Pay 1.50 now.", "Zahle 11.50 jetzt.", "en", "de"),
                         "Zahle 11.50 jetzt.")

    def test_idempotent(self):
        for src, tgt, s, t in [
                ("The amount is 12,450.00.", "Der Betrag ist 12,450.00.", "en", "de"),
                ("Znesek je 12.450,00 EUR.", "The amount is 12.450,00 EUR.", "sl", "en"),
                ("The price is CHF 12,450.50.", "Der Preis ist CHF 12,450.50.", "en", "de-CH")]:
            once = trlib.finish_draft(src, tgt, s, t)
            self.assertEqual(trlib.finish_draft(src, once, s, t), once)

    def test_swiss_finishing(self):
        self.assertEqual(trlib.finish_draft("The price is CHF 12,450.50.",
                                            "Der Preis ist CHF 12,450.50.", "en", "de-CH"),
                         f"Der Preis ist CHF 12{NBSP}450.50.")

    def test_grouped_whole_numbers(self):
        for src, tgt, s, t, want in [
                ("He holds 1,250 shares.", "Er hält 1,250 Aktien.", "en", "de",
                 "Er hält 1.250 Aktien."),
                ("He holds 1,250 shares.", "Ima 1,250 delnic.", "en", "sl",
                 "Ima 1.250 delnic."),
                ("Ima 1.250 delnic.", "He holds 1.250 shares.", "sl", "en",
                 "He holds 1,250 shares."),
                ("See item 1,250.", "Siehe Punkt 1,250.", "en", "de", "Siehe Punkt 1,250.")]:
            self.assertEqual(trlib.finish_draft(src, tgt, s, t), want, (src, t))

    def test_rewrites_of_ambiguous_numbers_are_listed(self):
        self.assertEqual(trlib.decimal_rewrites(
            "See Section 5.10 at 3.5 and EUR 2.50, 4.5% and 1,250.00.",
            "Siehe Abschnitt 5,10 zu 3,5 und EUR 2,50, 4,5 % und 1.250,00.", "en", "de"),
            ["5.10", "3.5"])
        self.assertEqual(trlib.decimal_rewrites("See 5.10.", "Siehe 5,10.", "sl", "en"), [])


class CompareNumbers(unittest.TestCase):
    """What tr-lint's NUM check and the invention retry compare."""

    def test_a_date_is_one_value_however_it_is_written(self):
        forms = ["5.3.2024", "5. 3. 2024", "05.03.2024", "2024-03-05", "March 5, 2024",
                 "March 5th, 2024", "5 March 2024", "5th of March 2024", "5. marec 2024",
                 "5. marca 2024", "5. März 2024"]
        for f in forms:
            self.assertEqual(trlib.norm_nums(f), trlib.norm_nums(forms[0]), f)

    def test_different_dates_differ(self):
        self.assertNotEqual(trlib.norm_nums("5.3.2024"), trlib.norm_nums("May 3, 2024"))

    def test_a_converted_date_adds_no_number(self):
        self.assertFalse(trlib.added_numbers(
            "Pogodba je bila sklenjena dne 5.3.2024 v Kranju.",
            "The contract was concluded on March 5, 2024 in Kranj."))
        self.assertFalse(trlib.added_numbers("Plačilo je zapadlo 15.04.2024.",
                                             "The payment fell due on April 15, 2024."))

    def test_may_as_a_verb_is_no_month(self):
        self.assertEqual(trlib.norm_nums("The court may order it in 2024."),
                         trlib.norm_nums("2024"))

    def gap(self, src, tgt, s, t):
        return any(trlib.compare_numbers(src, tgt, s, t))

    def test_numbers_are_read_in_each_sides_notation(self):
        self.assertTrue(self.gap("EUR 1.50", "EUR 150", "en", "de"))
        self.assertTrue(self.gap("12.5 m", "125 m", "en", "de"))
        self.assertTrue(self.gap("1,250 shares", "1,250 Aktien", "en", "de"))
        self.assertFalse(self.gap("12,450.00", "12.450,00", "en", "de"))
        self.assertFalse(self.gap("1,250 shares", "1.250 Aktien", "en", "de"))
        self.assertFalse(self.gap("3.25", "3,25", "en", "sl"))

    def test_conversions_the_kit_asks_for_add_nothing(self):
        for src, tgt, s, t in [
                ("The fee is CHF 20.", f"Die Gebühr beträgt Fr. 20.{WJ}–.", "en", "de-CH"),
                ("The fee is CHF 20.00.", "Die Gebühr beträgt CHF 20.", "en", "de-CH"),
                ("Odstavek 2 se lahko uporabi.", "Paragraph 2 may be applied.", "sl", "en"),
                ("The meeting is at 12:00 on Monday.",
                 "Die Sitzung ist um 12:00 am Montag.", "en", "de"),
                ("It starts at 2:30 p.m.", "Es beginnt um 14.30 Uhr.", "en", "de-CH"),
                ("The price is CHF 12,450.00.", "Der Preis beträgt CHF 12 450.00.",
                 "en", "de-CH"),
                ("1.- Introduction", "1. Einleitung", "en", "de"),
                ("Items 100 200 300", "Positionen 100, 200, 300", "en", "de"),
                ("The notice is dated 2024-03-05.",
                 "Die Kündigung datiert vom 5. März 2024.", "en", "de")]:
            self.assertEqual(trlib.compare_numbers(src, tgt, s, t),
                             (trlib.collections.Counter(), trlib.collections.Counter()),
                             (src, tgt))

    def test_a_slashed_date_agrees_with_either_reading(self):
        src = "The notice is dated 03/05/2024."
        for tgt in ("Die Kündigung datiert vom 05.03.2024.",
                    "Die Kündigung datiert vom 03.05.2024.",
                    "Die Kündigung datiert vom 5. März 2024."):
            self.assertFalse(self.gap(src, tgt, "en", "de"), tgt)
        self.assertTrue(self.gap(src, "Die Kündigung datiert vom 07.03.2024.", "en", "de"))
        self.assertFalse(self.gap("dated 03/15/2024", "vom 15.03.2024", "en", "de"))

    def test_reference_alignment_reads_each_side(self):
        src = ["The deposit is 1,250 EUR and is due on 03/05/2024."]
        tgt = ["Die Kaution beträgt 1.250 EUR und ist am 05.03.2024 fällig."]
        kept, _rejected = trref.align_document(src, tgt, langs=("en", "de"))
        self.assertEqual(kept, [(src[0], tgt[0], True)])
        kept, rejected = trref.align_document(["The rate is 12.5 per cent."],
                                              ["Der Satz beträgt 125 Prozent."],
                                              langs=("en", "de"))
        self.assertEqual((kept, [r[2] for r in rejected]), ([], ["numbers differ"]))


class RetrySelection(unittest.TestCase):
    """ollama_translate() against the mock: the first reply, the firmer retry,
    and which of them is kept."""

    @classmethod
    def setUpClass(cls):
        cls.mock = support.Mock()
        cls.ollama = trlib.OLLAMA
        trlib.OLLAMA = cls.mock.url

    @classmethod
    def tearDownClass(cls):
        trlib.OLLAMA = cls.ollama
        cls.mock.stop()

    def translate(self, src, first, firm):
        self.mock.set({"map": {src: first}, "map_firm": {src: firm}})
        before = len(self.mock.requests())
        out = trlib.ollama_translate(src, "sl", "en")
        return out, self.mock.requests()[before:]

    def test_a_converted_date_needs_no_retry(self):
        out, reqs = self.translate(
            "Pogodba je bila sklenjena dne 5.3.2024 v Kranju.",
            "The contract was concluded on March 5, 2024 in Kranj.",
            "The contract was concluded on 5.3.2024 in Kranj.")
        self.assertEqual(out, "The contract was concluded on March 5, 2024 in Kranj.")
        self.assertEqual([r["firm"] for r in reqs], [False])

    def test_the_reply_nearer_the_source_wins(self):
        first = "The payment of 300 EUR plus 45 EUR fell due on April 15, 2024."
        out, reqs = self.translate("Plačilo 300 EUR je zapadlo 15.04.2024.",
                                   first, "The payment fell due.")
        self.assertEqual(out, first)
        self.assertEqual([r["firm"] for r in reqs], [False, True])

    def test_the_firm_reply_wins_when_it_is_nearer(self):
        out, _reqs = self.translate("Številka zadeve je navedena spodaj.",
                                    "Case number 2 BvR 237/09 is given below.",
                                    "The case number is given below.")
        self.assertEqual(out, "The case number is given below.")

    def test_a_batch_retries_only_its_flagged_line(self):
        cells = [f"Opis blaga {c}{c}" for c in "abcdefghijklmnopq"] + [
            "Plačano dne 5.3.2024", "Številka spisa"]
        self.mock.set({"map": {"Plačano dne 5.3.2024": "Paid on March 5, 2024",
                               "Številka spisa": "File number 2 BvR 237/09"},
                       "map_firm": {"Številka spisa": "File number"}})
        before = len(self.mock.requests())
        out = trlib.ollama_translate_many(cells, "sl", "en")
        reqs = self.mock.requests()[before:]
        self.assertEqual([(r["batch"], r["firm"]) for r in reqs], [(True, False), (False, True)])
        self.assertEqual(out[0], "<<Opis blaga aa>>")
        self.assertEqual(out[-2:], ["Paid on March 5, 2024", "File number"])

    def test_a_value_with_letters_is_converted_in_code(self):
        before = len(self.mock.requests())
        self.assertEqual(trlib.ollama_translate("March 5, 2024", "en", "de"), "5. März 2024")
        self.assertEqual(trlib.ollama_translate_many(["2:30pm", "EUR 20"], "en", "de"),
                         ["14:30", "EUR 20"])
        self.assertEqual(len(self.mock.requests()), before)

    def test_a_reply_cut_off_is_retried_and_never_kept(self):
        src = "Najemnik plača najemnino do petega dne v mesecu."
        whole = "The tenant pays the rent by the fifth day of the month."
        self.mock.set({"map": {src: "The tenant pays the rent by the"}, "cut": [src],
                       "map_firm": {src: whole}})
        before = len(self.mock.requests())
        self.assertEqual(trlib.ollama_translate(src, "sl", "en"), whole)
        self.assertEqual([(r["firm"], r["cut"]) for r in self.mock.requests()[before:]],
                         [(False, True), (True, False)])
        cut = "Najemodajalec vzdržuje streho in fasado."
        self.mock.set({"map": {cut: "The landlord maintains"}, "cut": [cut], "cut_firm": [cut]})
        self.assertTrue(trlib.ollama_translate(cut, "sl", "en").startswith("[TRANSLATION FAILED]"))
        self.assertIsNone(trlib.tm_get(cut, "sl-en"))


class SwissDrafts(unittest.TestCase):
    """What finish_draft() makes of a Swiss draft written the way Germany
    writes, or left in English."""

    def finish(self, src, tgt):
        return trlib.finish_draft(src, tgt, "en", "de-CH")

    def test_germany_format_amounts_and_times(self):
        self.assertEqual(self.finish("The purchase price is CHF 12,450.00.",
                                     "Der Kaufpreis beträgt CHF 12.450,00."),
                         f"Der Kaufpreis beträgt Fr.{NBSP}12{NBSP}450.{WJ}–.")
        self.assertEqual(self.finish("The fee of EUR 1,250.50 is payable at 2:30 p.m.",
                                     "Die Gebühr von EUR 1.250,50 ist um 14:30 Uhr zu zahlen."),
                         "Die Gebühr von EUR 1250.50 ist um 14.30 Uhr zu zahlen.")

    def test_a_german_am_is_not_a_morning(self):
        self.assertEqual(self.finish("The meeting is at 12:00 on Monday.",
                                     "Die Sitzung ist um 12:00 am Montag."),
                         "Die Sitzung ist um 12.00 am Montag.")

    def test_a_range_shares_its_currency(self):
        self.assertEqual(self.finish("Between CHF 1,000.50 and 2,000.50",
                                     "Zwischen CHF 1,000.50 und 2,000.50"),
                         "Zwischen CHF 1000.50 und 2000.50")

    def test_grouped_whole_numbers(self):
        self.assertEqual(self.finish("He holds 12,450 shares.", "Er hält 12,450 Aktien."),
                         f"Er hält 12{NBSP}450 Aktien.")
        self.assertEqual(self.finish("He holds 1,250 shares.", "Er hält 1.250 Aktien."),
                         "Er hält 1250 Aktien.")

    def test_whole_francs_have_one_form(self):
        want = f"Fr.{NBSP}12{NBSP}450.{WJ}–"
        self.assertEqual(trlib.localize("12,450.00 CHF", "en", "de-CH"), want)
        for draft in ("Der Preis beträgt CHF 12,450.00.", "Der Preis beträgt 12,450.00 CHF."):
            self.assertEqual(self.finish("The price is CHF 12,450.00.", draft),
                             f"Der Preis beträgt {want}.", draft)
        for draft in ("Fr. 20.-", "CHF 20.–", "Fr. 20.—", f"Fr. 20.{WJ}–", "20.00 CHF"):
            self.assertEqual(self.finish("x", draft), f"Fr.{NBSP}20.{WJ}–", draft)
        self.assertEqual(self.finish("CHF 20.00 million", "CHF 20.00 Mio."), "CHF 20.00 Mio.")

    def test_whole_francs_keep_four_digits_together(self):
        self.assertEqual(self.finish("x", "Fr. 1 250.–"), f"Fr.{NBSP}1250.{WJ}–")
        self.assertEqual(self.finish("x", f"CHF 12{NBSP}450.00"), f"Fr.{NBSP}12{NBSP}450.{WJ}–")

    def test_fr_does_not_end_a_sentence(self):
        self.assertEqual(trlib.segment("Die Gebühr beträgt Fr. 20.– pro Monat."),
                         ["Die Gebühr beträgt Fr. 20.– pro Monat."])

    def test_sharp_s_is_kept_in_any_case_and_inflected(self):
        self.assertEqual(trlib.swiss_spelling("STRAßE 5", "Straße 5"), "Straße 5")
        self.assertEqual(trlib.swiss_spelling("Mr Strauß", "Straußens Vertrag"),
                         "Straußens Vertrag")
        self.assertEqual(trlib.swiss_spelling("The street", "Die Straße"), "Die Strasse")


class SwissLint(unittest.TestCase):
    """tr-run and tr-lint on a Swiss project whose model wrote the forms the
    Swiss prompt asks for, and Germany's."""

    @classmethod
    def setUpClass(cls):
        cls.mock = support.Mock()

    @classmethod
    def tearDownClass(cls):
        cls.mock.stop()

    def test_the_kits_own_swiss_forms_raise_nothing(self):
        root = support.project("swiss-lint", "en", "de-CH")
        src, out = f"{root}/source/fees.docx", f"{root}/translated/fees.docx"
        support.write_docx(src, ["The fee is CHF 20.", "The price is 12,450.00 CHF in total.",
                                 "The hearing is at 14:30."])
        self.mock.set({"map": {
            "The fee is CHF 20.": "Die Gebühr beträgt Fr. 20.–.",
            "The price is 12,450.00 CHF in total.": "Der Preis beträgt 12.450,00 CHF insgesamt.",
            "The hearing is at 14:30.": "Die Verhandlung ist um 14:30 Uhr."}})
        code, log = support.run("tr-run", root, src, mock=self.mock)
        self.assertEqual(code, 0, log)
        self.assertEqual(support.read_docx(out), [
            f"Die Gebühr beträgt Fr.{NBSP}20.{WJ}–.",
            f"Der Preis beträgt Fr.{NBSP}12{NBSP}450.{WJ}– insgesamt.",
            "Die Verhandlung ist um 14.30 Uhr."])
        _code, lint = support.run("tr-lint", root)
        self.assertRegex(lint, r"(?m)^findings:\s+0$")


class TypedCells(unittest.TestCase):
    """tr-xlsx: a typed date cell takes the target's date format; a number, or
    a date with a time of day, keeps its own. Audit 2026-09-10: P-3."""

    def test_a_typed_date_takes_the_targets_format(self):
        import datetime
        import os
        import openpyxl
        root = support.project("typed-cells", "sl", "en")
        src, out = f"{root}/source/t.xlsx", f"{root}/translated/t.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"] = (datetime.date(2024, 3, 5),
                                        datetime.datetime(2024, 3, 5, 14, 30), 12450.5)
        ws["A1"].number_format = "d.m.yyyy"
        ws["B1"].number_format = "d.m.yyyy h:mm"
        ws["C1"].number_format = "#,##0.00"
        os.makedirs(os.path.dirname(src), exist_ok=True)
        wb.save(src)
        mock = support.Mock()
        try:
            code, log = support.run("tr-xlsx", root, src, out, mock=mock)
        finally:
            mock.stop()
        self.assertEqual(code, 0, log)
        ws = openpyxl.load_workbook(out).active
        self.assertEqual([ws[c].number_format for c in ("A1", "B1", "C1")],
                         ["[$-409]mmmm d, yyyy", "d.m.yyyy h:mm", "#,##0.00"])


if __name__ == "__main__":
    unittest.main()
