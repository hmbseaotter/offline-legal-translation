"""Shared set-up for the kit's tests: a throwaway project root, invented
documents, and a stand-in for Ollama.

Every test module imports this before trlib, because trlib reads TR_PROJECTS,
TR_ROOT and TR_OLLAMA when it is imported. Nothing here reads or writes under
~/translation-work: the root is a new temporary directory, refused if it
would be anywhere under there, and removed when the tests end unless
TR_TESTS_KEEP is set. TR_OLLAMA points at a closed port unless a test starts
the mock, so no test can reach a real model by accident.
"""
import atexit
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BIN = os.path.join(KIT, "bin")

ROOT = tempfile.mkdtemp(prefix="tr-tests-")
if os.path.realpath(ROOT).startswith(
        os.path.realpath(os.path.expanduser("~/translation-work")) + os.sep):
    sys.exit(f"refusing to test under ~/translation-work: {ROOT}")
if os.environ.get("TR_TESTS_KEEP"):
    print(f"test root kept: {ROOT}", file=sys.stderr)
else:
    atexit.register(shutil.rmtree, ROOT, True)

for _k in ("TR_MODEL", "TR_SRC", "TR_TGT", "TR_PROMPT_VERSION", "TR_SUFFIX",
           "TR_OCR_LANGS"):
    os.environ.pop(_k, None)
os.environ.update(TR_PROJECTS=ROOT, TR_ROOT=os.path.join(ROOT, "library"),
                  TR_OLLAMA="http://127.0.0.1:9", PYTHONDONTWRITEBYTECODE="1")
os.makedirs(os.path.join(ROOT, "library", "work"), exist_ok=True)
sys.path.insert(0, os.path.join(KIT, "lib"))

MODELS = ["gams3:q8", "eurollm9b-2512:q8"]


def project(name, src, tgt):
    """A project directory under the test root, with its pair in project.conf."""
    root = os.path.join(ROOT, name)
    for d in ("source", "translated", "work/ocr", "logs", "glossary", "prompts",
              "reference"):
        os.makedirs(os.path.join(root, d), exist_ok=True)
    with open(os.path.join(root, "project.conf"), "w", encoding="utf-8") as fh:
        fh.write(f"TR_SRC={src}\nTR_TGT={tgt}\n")
    return root


def run(tool, root, *args, mock=None):
    """(exit status, stdout and stderr) of bin/<tool> against a project."""
    env = dict(os.environ, TR_ROOT=root)
    if mock:
        env["TR_OLLAMA"] = mock.url
    r = subprocess.run([os.path.join(BIN, tool), *args], env=env,
                       stdin=subprocess.DEVNULL, capture_output=True, text=True,
                       timeout=900)
    return r.returncode, r.stdout + r.stderr


def write_docx(path, paragraphs):
    import docx
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(path)


def read_docx(path):
    import docx
    return [p.text for p in docx.Document(path).paragraphs if p.text.strip()]


def write_text(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


def write_pdf(path, lines):
    """A one-page PDF with a real text layer, one line per entry: what a word
    processor exports, so tr-pdf takes it as born-digital. Helvetica in
    WinAnsi, which holds German's letters."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    y, ops = 800, []
    for ln in lines:
        ops.append(f"BT /F1 10 Tf 50 {y} Td ({esc(ln)}) Tj ET")
        y -= 28
    stream = "\n".join(ops).encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R"
            b" /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica"
            b" /Encoding /WinAnsiEncoding >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objs) + 1, xref)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(out)


class Mock:
    """tests/mock_ollama.py on a free port, answering by rules (see there)."""

    def __init__(self, rules=None):
        self.dir = tempfile.mkdtemp(prefix="mock-", dir=ROOT)
        self.rules = os.path.join(self.dir, "rules.json")
        self.log = os.path.join(self.dir, "log.jsonl")
        self.set(rules or {})
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            port = s.getsockname()[1]
        self.url = f"http://127.0.0.1:{port}"
        self.proc = subprocess.Popen(
            [sys.executable, os.path.join(KIT, "tests", "mock_ollama.py"),
             str(port), self.rules, self.log])
        atexit.register(self.stop)
        for _ in range(200):
            try:
                urllib.request.urlopen(self.url + "/api/tags", timeout=1).close()
                return
            except OSError:
                time.sleep(0.05)
        raise RuntimeError("mock Ollama did not start")

    def set(self, rules):
        with open(self.rules, "w", encoding="utf-8") as fh:
            json.dump({"models": MODELS, "wrap": "<<{}>>", **rules}, fh,
                      ensure_ascii=False)

    def requests(self):
        if not os.path.exists(self.log):
            return []
        with open(self.log, encoding="utf-8") as fh:
            return [json.loads(ln) for ln in fh if ln.strip()]

    def stop(self):
        if self.proc.poll() is None:
            self.proc.terminate()
            self.proc.wait(10)
